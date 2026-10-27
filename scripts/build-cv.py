from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import argparse
import calendar
import json

root = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description='Build the reviewed CV source from the shared profile. Render and inspect before publishing the PDF.')
parser.add_argument('--profile', type=Path, default=root / 'content/profile.json')
parser.add_argument('--output', type=Path, default=root / 'output/docx/Letlhogonolo-Fanampe-AI-Solutions-Engineer-CV.docx')
args = parser.parse_args()
profile = json.loads(args.profile.read_text())
roles = {item['id']: item for item in profile['experience']}
projects = {item['id']: item for item in profile['projects']}
args.output.parent.mkdir(parents=True, exist_ok=True)

def plain(text):
 return text.replace('—', '-').replace('–', '-').replace('‑', '-').replace('’', "'")

def month(value):
 if not value: return 'Present'
 year, number = value.split('-')
 return f'{calendar.month_abbr[int(number)]} {year}'

def role_dates(item):
 return f"{month(item['start'])} - {month(item['end'])}"

doc=Document(); s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.6);s.bottom_margin=Inches(.6);s.left_margin=Inches(.7);s.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(4)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.04
for name,size in [('Title',24),('Subtitle',14),('Heading 1',12),('Heading 2',11)]:
 st=doc.styles[name];st.font.size=Pt(size);st.paragraph_format.space_before=Pt(9 if 'Heading' in name else 0)
 st.paragraph_format.keep_with_next=True
for name in ['Heading 1','Heading 2']:doc.styles[name].font.bold=True
st=doc.styles['List Bullet'];st.font.size=Pt(11);st.paragraph_format.space_after=Pt(3);st.paragraph_format.left_indent=Inches(.14);st.paragraph_format.first_line_indent=Inches(-.14)
def para(t='',style=None):return doc.add_paragraph(plain(t),style)
def bullet(t):return para(t,'List Bullet')
def section(t):return para(t,'Heading 1')
def role(role_id):
 item=roles[role_id]
 p=para(item['title'],'Heading 2');p.paragraph_format.space_after=Pt(2)
 context=' | '.join(str(value) for value in [item['organization'],item.get('type'),item.get('location'),role_dates(item)] if value)
 p=para(context);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4);p.runs[0].font.size=Pt(10)

def link(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'174E63');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
para(profile['name'],'Title');para(profile['role'],'Subtitle')
p=para(profile['location']+' | Remote and hybrid | Available immediately');p.runs[0].font.size=Pt(10)
p=para(profile['contact']['email']+' | '+profile['contact']['phone']);p.runs[0].font.size=Pt(10)
p=para();link(p,profile['contact']['website'].removeprefix('https://'),profile['contact']['website']);p.add_run(' | ');link(p,'LinkedIn',profile['linkedin']);p.add_run(' | ');link(p,'GitHub',profile['contact']['github'])
section('Profile')
para('AI Solutions Engineer building LLM applications that connect business knowledge, data and workflows. Hands-on experience delivering RAG and SQL assistants, API integrations, OCR pipelines and Linux deployments. Combines application engineering with AI response evaluation and human feedback for model training.')
section('Core skills')
para('Python, SQL, FastAPI, LangChain, Flask, Django, Streamlit, REST APIs, OpenAI APIs, Together AI, RAG, NLP, LLM fine-tuning, FAISS, DeepLake, ChromaDB, SQLite, Docker, Git, Linux, Nginx, systemd, OCR, RLHF and model evaluation.')
section('Professional experience')
freelance=roles['freelance']
concurrent=[item['referenceName'] for item in profile['experience'] if item['id']!='freelance' and 'Contract' in item.get('type','') and max(item['start'],freelance['start'])<=min(item.get('end') or '9999-12',freelance.get('end') or '9999-12')]
if concurrent:
 names=', '.join(concurrent[:-1])+' and '+concurrent[-1] if len(concurrent)>1 else concurrent[0]
 p=para(f'The {names} contracts overlapped with freelance work.');p.paragraph_format.space_after=Pt(4)

role('freelance')
bullet('Build AI applications and developer tools using Python, language models and external APIs, from early prototypes through integration and demonstrations.')
bullet('Developed VocalThread for YouTube comment analysis and speech generation, and built local recovery and agent workflow tools described in the selected projects below.')
role('kwantu')
bullet('Developed and deployed RAG assistants for PDF documents and SQL databases with Python, LangChain, OpenAI embeddings, FAISS and DeepLake, including multilingual retrieval.')
bullet('Integrated chatbot services with FastAPI, WhatsApp Business API and Telegram using models and APIs including Mixtral and Together AI.')
bullet('Built an ID-document processing pipeline covering data cleaning, Label Studio annotation, Detectron2 training and OCR with Tesseract and PaddleOCR.')
bullet('Maintained CentOS/Linux deployments with systemd and Nginx; resolved SQL schema mismatches, API rate limits, connectivity problems and inconsistent model responses.')
role('outlier')
bullet('Evaluated and ranked model responses for accuracy, coherence, relevance and adherence to guidelines; verified training data and reviewed peer submissions.')
bullet('Designed prompts and adversarial tests to identify reasoning errors and model failure cases, and documented feedback for reinforcement learning from human feedback.')
doc.add_page_break()
section('Professional experience continued')
role('translated')
bullet('Fact-checked, rated and rewrote AI-generated prompts and responses for a major GPU enterprise client, documenting accuracy, relevance and instruction-following issues.')
role('next-sapien')
bullet(roles['next-sapien']['bullets'][0])
section('Selected AI projects')
def project(title,url,text):
 p=para(title,'Heading 2');p.paragraph_format.space_after=Pt(3)
 p=para(text);p.add_run(' ');link(p,'Project',url);p.paragraph_format.space_after=Pt(5)
project('Agent Handoff Kit','https://djpapzin.com/case-studies/agent-handoff','Built a Python and SQLite prototype that resumes interrupted local work from saved checkpoints. Uses transactional receipts, leases and generation fencing to suppress duplicate effects within a local database transaction. Published a recovery demo, architecture notes, source code and tests.')
project('AgentOps Mobile Command Center','https://github.com/djpapzin/agentops-mobile-command-center','Built a FastAPI and SQLite control room for goals, approvals, run history and model routing through Telegram. Includes Docker setup, a demonstration mode and optional integrations with model providers.')
project('TruthGuard','https://github.com/djpapzin/TruthGuard-AI-Fake-News-Detection-with-LLM',projects['truthguard']['summary'])
project('VocalThread','https://github.com/djpapzin/Vocal-Thread','Developed a Streamlit application that analyses YouTube comments and turns them into spoken dialogue using ElevenLabs, YouTube Data API and Google Gemini. Features include sentiment analysis and chat with comments.')
section('Earlier experience')
role('afrisam')
para('Performed routine and specialised laboratory testing, maintained analytical equipment and supported product quality assurance. Built a foundation in data analysis, troubleshooting and consistent evaluation procedures.')
section('Education and professional development')
para('National Diploma in Analytical Chemistry | Tshwane University of Technology | 2010 - 2015')
para('LangChain and Vector Databases in Production | Activeloop | Sep 2023')
para('Master Python 3 | Udemy | Aug 2023')
para('Career Essentials in System Administration | Microsoft and LinkedIn | Mar 2023')
section('Recognition')
para('WebGPU Hackathon One-Day Challenge winner with VisualPro | Aug 2023. Built Node.js API endpoints and SQLite storage for a team 3D data visualisation application.')
for el in doc.styles.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
for el in doc._element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
doc.core_properties.title='Letlhogonolo Fanampe AI Solutions Engineer CV'
doc.core_properties.author='Letlhogonolo Fanampe'
doc.core_properties.subject='AI Solutions Engineer experience and selected projects'
doc.save(args.output)
print(f'Saved {args.output}')
